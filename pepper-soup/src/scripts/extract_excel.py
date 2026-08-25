from __future__ import annotations

import hashlib
import json
import logging
import re
import sys
from datetime import datetime, date
from typing import Any

import openpyxl
from openpyxl.worksheet.worksheet import Worksheet

# Logging to stderr — does not pollute the JSON stdout output.
logging.basicConfig(stream=sys.stderr, level=logging.INFO, format="[excel] %(message)s")
_log = logging.getLogger(__name__)


def cell_str(v: Any) -> str:
    if v is None:
        return ""
    if isinstance(v, (datetime, date)):
        return v.strftime("%Y-%m-%d")
    return re.sub(r"\s{2,}", " ", re.sub(r"\n+", " | ", str(v).strip())).strip()


def non_empty_rows(ws: Worksheet) -> list[tuple[int, list[Any]]]:
    return [
        (i, list(row))
        for i, row in enumerate(ws.iter_rows(values_only=True), 1)
        if any(c is not None for c in row)
    ]


def compact(row: list[Any]) -> list[str]:
    return [cell_str(c) for c in row if c is not None and str(c).strip()]


def sanitize(s: str) -> str:
    return re.sub(r"[^a-zA-Z0-9]+", "_", s).strip("_")


def source_id(stem: str, sheet: str) -> str:
    return f"{sanitize(stem)}__{sanitize(sheet)}"


TOPIC_RULES: list[tuple[str, str]] = [
    (
        r"\b(gl\b|debit|credit|\bdr\b|\bcr\b|ledger|loan receivable|deferred income|liability|asset type)\b",
        "accounting",
    ),
    (
        r"\b(charge.?off|write.?off|written.?off|preclos|reversal|overpaid|cbr|bad.?debt)\b",
        "edge_case",
    ),
    (
        r"\b(to be discussed|not yet|tbd|pp yet|unclear|pending|new bucket)\b",
        "open_question",
    ),
    (
        r"\b(validat|must not|should not|cannot|constraint|maximum|minimum)\b",
        "constraint",
    ),
    (
        r"\b(interest rate|days in year|days in month|repaid every|product level|parameter)\b",
        "configuration",
    ),
    (r"\b(amortiz|cob|daily|accrual|repayment|payoff|preclosure)\b", "behavior"),
    (
        r"\b(capitalized income|origination fee|deferred fee|loan portfolio|fee bucket)\b",
        "concept",
    ),
    (r"\b(event|trigger|webhook)\b", "business_event"),
    (r"\b(api|endpoint|\bhttp\b|\bjson\b)\b", "api"),
    (r"\b(table|column|schema|database)\b", "database"),
    (r"\b(integrat|external system|third.?party)\b", "integration"),
]

# Pre-compiled patterns for hot-path lookups (avoid recompiling on every call).
_TOPIC_PATS = [(re.compile(p, re.I), t) for p, t in TOPIC_RULES]
_STATUS_OPEN = re.compile(
    r"\b(to be discussed|not yet|tbd|pp yet|open|unclear|pending|new bucket)\b", re.I
)
_STATUS_INFERRED = re.compile(r"\b(implied|assumed|inferred|should|likely)\b", re.I)
# Single regex covering all recognized parameter key names for params_block().
_PARAMS_RE = re.compile(
    r"interest rate|days in year|days in month|disbursement|origination fees"
    r"|capitalized income|total amount|repaid every|capitalized income date"
    r"|maturity date|income amount|capitalized income adjustment date"
    r"|capitalized income adjustment amount|should iterate|iterate with",
    re.I,
)
_UC_PAT = re.compile(r"^uc[\s\-_]?\d+", re.I)


def _get(cols: list[str], c: int | None) -> str:
    """Return stripped cell value at index *c*, or '' if None or out-of-range."""
    return cols[c].strip() if c is not None and c < len(cols) else ""


def topic(text: str) -> str:
    lo = text.lower()
    for pat, t in _TOPIC_PATS:
        if pat.search(lo):
            return t
    return "other"


def status(text: str) -> str:
    if _STATUS_OPEN.search(text):
        return "open"
    if _STATUS_INFERRED.search(text):
        return "inferred"
    return "confirmed"


def confidence(s: str) -> str:
    return {"open": "low", "inferred": "medium"}.get(s, "high")


def fact(
    sid: str,
    fid: str,
    idx: int,
    label: str,
    content: str,
    t: str | None = None,
    s: str | None = None,
    c: str | None = None,
) -> dict:
    full = f"{label} {content}"
    rs = s or status(full)
    return {
        "id": f"fact_{hashlib.md5(f'{sid}_{idx}'.encode()).hexdigest()[:12]}",
        "topic": t or topic(full),
        "label": label[:200],
        "content": content[:2000],
        "confidence": c or confidence(rs),
        "status": rs,
        "provenance": {"source_type": "excel", "source_id": sid, "feature_id": fid},
    }


def example(
    sid: str,
    fid: str,
    idx: int,
    title: str,
    scenario: str,
    steps: list[dict],
    result: str,
) -> dict:
    return {
        "id": f"ex_{hashlib.md5(f'{sid}_{idx}'.encode()).hexdigest()[:12]}",
        "title": title[:200],
        "scenario": scenario[:500],
        "steps": steps,
        "result": result[:1000],
        "provenance": {"source_type": "excel", "source_id": sid, "feature_id": fid},
    }


def find_header(rows: list[tuple[int, list]], keywords: list[str]) -> int | None:
    kws = [k.lower() for k in keywords]
    for i, (_, row) in enumerate(rows):
        txt = " ".join(cell_str(c).lower() for c in row if c)
        if all(k in txt for k in kws):
            return i
    return None


def params_block(rows: list[tuple[int, list]], max_rows: int = 30) -> dict[str, str]:
    out: dict[str, str] = {}
    for _, row in rows[:max_rows]:
        cols = compact(row)
        if len(cols) >= 2 and _PARAMS_RE.search(cols[0]):
            out[cols[0]] = " | ".join(cols[1:])
    return out


def parse_schedule(rows: list[tuple[int, list]], hi: int) -> list[dict]:
    hv = [cell_str(c).lower() for c in rows[hi][1]]

    def ci(n: str) -> int | None:
        return next((j for j, h in enumerate(hv) if n in h), None)

    dc, bc, ec, pc, ic = ci("date"), ci("balance"), ci("emi"), ci("prin"), ci("int")
    steps, n = [], 1
    for _, row in rows[hi + 1 :]:
        cols = [cell_str(c) for c in row]
        dv = _get(cols, dc)
        if not dv or not re.match(r"\d{4}-\d{2}-\d{2}", dv):
            continue
        if _get(cols, 0).lower() in ("total", ""):
            break
        pts = [f"Date: {dv}"]
        if _get(cols, bc):
            pts.append(f"Balance: ${_get(cols, bc)}")
        if _get(cols, ec):
            pts.append(f"EMI: ${_get(cols, ec)}")
        if _get(cols, pc):
            pts.append(f"Principal: ${_get(cols, pc)}")
        if _get(cols, ic):
            pts.append(f"Interest: ${_get(cols, ic)}")
        steps.append({"step": n, "description": " | ".join(pts)})
        n += 1
    return steps


def parse_txn_table(
    rows: list[tuple[int, list]], hi: int, sid: str, fid: str, fo: int
) -> tuple[list[dict], list[dict]]:
    hv = [cell_str(c).lower() for c in rows[hi][1]]

    def ci(n: str) -> int | None:
        return next((j for j, h in enumerate(hv) if n in h), None)

    tc, dc = ci("transaction type"), ci("transaction date") or ci("date")
    dbc, crc = ci("db gl") or ci(" db"), ci("cr gl") or ci(" cr")
    facts_out, steps, n = [], [], 1
    skip = {"accrual", "transaction type"}
    for _, row in rows[hi + 1 :]:
        cols = [cell_str(c) for c in row]
        txn = _get(cols, tc)
        if not txn:
            if (_get(cols, dbc) or _get(cols, crc)) and steps:
                suf = []
                if _get(cols, dbc):
                    suf.append(f"DR: {_get(cols, dbc)[:80]}")
                if _get(cols, crc):
                    suf.append(f"CR: {_get(cols, crc)[:80]}")
                steps[-1]["description"] += " | " + " | ".join(suf)
            continue
        amts = ", ".join(
            f"${c}" for c in cols if re.match(r"^-?[\d]+\.?\d*$", c.strip())
        )[:60]
        pts = [f"[{txn}]"]
        if _get(cols, dc):
            pts.append(f"date={_get(cols, dc)}")
        if amts:
            pts.append(amts)
        if _get(cols, dbc):
            pts.append(f"DR: {_get(cols, dbc)[:70]}")
        if _get(cols, crc):
            pts.append(f"CR: {_get(cols, crc)[:70]}")
        steps.append({"step": n, "description": " | ".join(pts)})
        n += 1
        if txn.lower() not in skip and (_get(cols, dbc) or _get(cols, crc)):
            label = f"Transaction: {txn}" + (f" ({_get(cols, dc)})" if _get(cols, dc) else "")
            content = " | ".join(
                filter(
                    None,
                    [
                        f"Type: {txn}",
                        f"Date: {_get(cols, dc)}" if _get(cols, dc) else "",
                        f"Amounts: {amts}" if amts else "",
                        f"DEBIT GL: {_get(cols, dbc)}" if _get(cols, dbc) else "",
                        f"CREDIT GL: {_get(cols, crc)}" if _get(cols, crc) else "",
                    ],
                )
            )
            facts_out.append(
                fact(
                    sid,
                    fid,
                    fo + len(facts_out),
                    label,
                    content,
                    t="accounting",
                    s="confirmed",
                    c="high",
                )
            )
    return facts_out, steps


def sheet_type(ws: Worksheet, name: str) -> str:
    n = name.lower().strip()
    if n == "questions":
        return "qa"
    if "accounting" in n:
        return "accounting"
    if n == "acs":
        return "acs"
    if _UC_PAT.match(n):
        return "use_case"
    if n in ("example calculation", "schedule example"):
        return "schedule_example"
    if "amortization logic" in n:
        return "amortization"
    if "playground" in n:
        return "playground"
    return "generic"


def parse_qa(ws, sid, fid):
    facts_out, i = [], 0
    OPEN = {
        "to be discussed",
        "pp yet to share",
        "new bucket?",
        "tbd",
        "not yet",
        "later",
    }
    for _, row in non_empty_rows(ws):
        cols = compact(row)
        if len(cols) < 1 or cols[0].lower() in ("questions", "answers"):
            continue
        q = cols[0]
        a = cols[1] if len(cols) > 1 else ""
        is_open = not a or any(s in a.lower() for s in OPEN)
        label = f"Open Question: {q[:120]}"
        content = f"Question: {q}" + (f" | Answer: {a}" if a else "")
        facts_out.append(
            fact(
                sid,
                fid,
                i,
                label,
                content,
                t="open_question",
                s="open" if is_open else "confirmed",
                c="low" if is_open else "high",
            )
        )
        i += 1
    return facts_out, []


def parse_accounting(ws, sid, fid):
    facts_out, i = [], 0
    for _, row in non_empty_rows(ws):
        cols = compact(row)
        if len(cols) < 3 or cols[0].lower() in (
            "new transaction type",
            "transaction type",
        ):
            continue
        txn, dr, cr = cols[0], cols[1], cols[2]
        facts_out.append(
            fact(
                sid,
                fid,
                i,
                label=f"Accounting Entry: {txn}",
                content=f"Transaction '{txn}' — DEBIT: {dr} | CREDIT: {cr}",
                t="accounting",
                s="confirmed",
                c="high",
            )
        )
        i += 1
    return facts_out, []


def parse_acs(ws, sid, fid):
    facts_out = []
    for i, (_, row) in enumerate(non_empty_rows(ws)):
        cols = compact(row)
        if not cols:
            continue
        text = " | ".join(cols)
        if len(text) < 5:
            continue
        facts_out.append(fact(sid, fid, i, f"AC: {text[:120]}", text))
    return facts_out, []


def parse_uc(ws, sid, fid, name):
    facts_out, examples_out = [], []
    rows = non_empty_rows(ws)
    if not rows:
        return facts_out, examples_out

    title_cols = compact(rows[0][1])
    title = title_cols[0] if title_cols else name
    p = params_block(rows)
    fi = 0

    for k, v in p.items():
        facts_out.append(
            fact(
                sid,
                fid,
                fi,
                f"Loan Parameter: {k}",
                f"Scenario '{title}': {k} = {v}",
                t="configuration",
                s="confirmed",
                c="high",
            )
        )
        fi += 1

    facts_out.append(
        fact(
            sid,
            fid,
            fi,
            f"Scenario: {title}",
            f"Use case '{title}'. Params: {' | '.join(f'{k}={v}' for k,v in list(p.items())[:6])}",
            t="concept",
            s="confirmed",
            c="high",
        )
    )
    fi += 1

    sched_hi = find_header(rows, ["date", "balance", "emi"])
    repay_steps = parse_schedule(rows, sched_hi) if sched_hi is not None else []

    txn_hi = find_header(rows, ["transaction type", "transaction date"])
    txn_facts, txn_steps = ([], [])
    if txn_hi is not None:
        txn_facts, txn_steps = parse_txn_table(rows, txn_hi, sid, fid, fi)
        facts_out.extend(txn_facts)
        fi += len(txn_facts)

    all_steps, n = [], 1
    for s in repay_steps:
        all_steps.append({"step": n, "description": "[Schedule] " + s["description"]})
        n += 1
    for s in txn_steps:
        all_steps.append({"step": n, "description": "[Txn/GL] " + s["description"]})
        n += 1

    if all_steps or p:
        scenario_str = " | ".join(f"{k}: {v}" for k, v in list(p.items())[:5])
        examples_out.append(
            example(
                sid,
                fid,
                0,
                title,
                scenario_str or title,
                all_steps,
                f"Transactions and GL entries applied per '{title}'.",
            )
        )

    return facts_out, examples_out


def parse_schedule_example(ws, sid, fid):
    facts_out, examples_out = [], []
    rows = non_empty_rows(ws)

    uc_starts = [
        i
        for i, (_, row) in enumerate(rows)
        if re.match(r"^UC-?\d+", cell_str(row[0]).strip(), re.I)
    ]

    for bn, start in enumerate(uc_starts):
        end = uc_starts[bn + 1] if bn + 1 < len(uc_starts) else len(rows)
        block = rows[start:end]
        tc = compact(block[0][1])
        title = tc[0] if tc else f"Block {bn + 1}"
        p = params_block(block)
        for i, (k, v) in enumerate(p.items()):
            facts_out.append(
                fact(
                    sid,
                    fid,
                    len(facts_out),
                    f"Loan Parameter ({title}): {k}",
                    f"Scenario '{title}': {k} = {v}",
                    t="configuration",
                    s="confirmed",
                    c="high",
                )
            )
        hi = find_header(block, ["date", "balance", "emi"])
        steps = parse_schedule(block, hi) if hi is not None else []
        if steps:
            scenario_str = " | ".join(f"{k}: {v}" for k, v in list(p.items())[:4])
            examples_out.append(
                example(
                    sid,
                    fid,
                    bn,
                    title,
                    scenario_str or title,
                    steps,
                    f"Repayment schedule for '{title}'.",
                )
            )
    return facts_out, examples_out


def parse_amortization(ws, sid, fid, name):
    facts_out, examples_out = [], []
    rows = non_empty_rows(ws)
    if not rows:
        return facts_out, examples_out

    p: dict[str, str] = {}
    for _, row in rows[:8]:
        cols = compact(row)
        if len(cols) >= 2:
            p[cols[0]] = " | ".join(cols[1:])

    for i, (k, v) in enumerate(p.items()):
        facts_out.append(
            fact(
                sid,
                fid,
                i,
                f"Amortization Parameter: {k}",
                f"Sheet '{name}': {k} = {v}",
                t="configuration",
                s="confirmed",
                c="high",
            )
        )

    hi = find_header(rows, ["transaction type", "transaction date"])
    amort_steps, n = [], 1
    if hi is not None:
        hv = [cell_str(c).lower() for c in rows[hi][1]]

        def ci(nm):
            return next((j for j, h in enumerate(hv) if nm in h), None)

        tc, dc, ec, fc = (
            ci("transaction type"),
            ci("transaction date"),
            ci("emi"),
            ci("fees"),
        )
        for _, row in rows[hi + 1 :]:
            cols = [cell_str(c) for c in row]
            txn = _get(cols, tc)
            if not txn:
                continue
            pts = [f"[{txn}]"]
            if _get(cols, dc):
                pts.append(f"date={_get(cols, dc)}")
            if _get(cols, ec):
                pts.append(f"daily_amort=${_get(cols, ec)}")
            if _get(cols, fc):
                pts.append(f"fees_bal=${_get(cols, fc)}")
            amort_steps.append({"step": n, "description": " | ".join(pts)})
            n += 1

    if amort_steps:
        scenario_str = " | ".join(f"{k}: {v}" for k, v in list(p.items())[:4])
        adj = p.get("Capitalized Income Adjustment Date", "")
        adj_amt = p.get("Capitalized Income Adjustment Amount", "")
        result = (
            f"Daily COB amortization for {len(amort_steps)} day(s). "
            "DR Deferred Income (Liability), CR Income from Amortization (Income)."
        )
        if adj and adj_amt:
            result += f" Mid-term adjustment of {adj_amt} on {adj}; balance re-amortized over remaining days."
        examples_out.append(
            example(
                sid,
                fid,
                0,
                f"Daily Amortization Schedule: {name}",
                f"Parameters: {scenario_str}",
                amort_steps,
                result,
            )
        )

    income = p.get("Income Amount", "")
    cap_dt = p.get("Capitalized Income date", "")
    if income or cap_dt:
        content = (
            f"Capitalized income of {income} originated on {cap_dt}, "
            f"maturing on {p.get('Maturity Date', 'N/A')}. "
            "Amortized daily via COB: DR Deferred Income GL, CR Income from Amortization GL."
        )
        facts_out.append(
            fact(
                sid,
                fid,
                len(facts_out),
                f"Amortization Behavior ({name})",
                content,
                t="behavior",
                s="confirmed",
                c="high",
            )
        )
    return facts_out, examples_out


def parse_playground(ws, sid, fid):
    facts_out, examples_out = [], []
    rows = non_empty_rows(ws)
    if not rows:
        return facts_out, examples_out

    tc = compact(rows[0][1])
    title = tc[0] if tc else "Amortization Job Playground"
    p = params_block(rows)
    for i, (k, v) in enumerate(p.items()):
        facts_out.append(
            fact(
                sid,
                fid,
                i,
                f"Playground Parameter: {k}",
                f"'{title}': {k} = {v}",
                t="configuration",
                s="confirmed",
                c="high",
            )
        )

    hi = find_header(rows, ["transaction type", "transaction date"])
    if hi is not None:
        txn_facts, steps = parse_txn_table(rows, hi, sid, fid, len(facts_out))
        facts_out.extend(txn_facts)
        if steps:
            scenario_str = " | ".join(f"{k}: {v}" for k, v in list(p.items())[:4])
            examples_out.append(
                example(
                    sid,
                    fid,
                    0,
                    title,
                    f"Parameters: {scenario_str}. COB executed day by day.",
                    steps,
                    "Each COB cycle: DR Deferred Income GL, CR Income from Amortization GL.",
                )
            )
    return facts_out, examples_out


def parse_generic(ws, sid, fid, name):
    facts = [
        fact(
            sid,
            fid,
            i,
            f"{name} R{rn}: {(' | '.join(compact(row)))[:80]}",
            " | ".join(compact(row)),
        )
        for i, (rn, row) in enumerate(non_empty_rows(ws))
        if len(" ".join(compact(row))) >= 5
    ]
    return facts, []


def process_sheet(ws: Worksheet, name: str, stem: str, fid: str) -> dict:
    sid = source_id(stem, name)
    t = sheet_type(ws, name)
    _log.info("Sheet '%s' → type=%s  sid=%s", name, t, sid)
    dispatch = {
        "qa": lambda: parse_qa(ws, sid, fid),
        "accounting": lambda: parse_accounting(ws, sid, fid),
        "acs": lambda: parse_acs(ws, sid, fid),
        "use_case": lambda: parse_uc(ws, sid, fid, name),
        "schedule_example": lambda: parse_schedule_example(ws, sid, fid),
        "amortization": lambda: parse_amortization(ws, sid, fid, name),
        "playground": lambda: parse_playground(ws, sid, fid),
        "generic": lambda: parse_generic(ws, sid, fid, name),
    }
    facts_out, examples_out = dispatch.get(t, dispatch["generic"])()
    _log.info("  → %d fact(s), %d example(s)", len(facts_out), len(examples_out))
    return {
        "meta": {
            "source_type": "excel",
            "source_id": sid,
            "content_type": "scenario" if examples_out else "table",
            "feature_id": fid,
        },
        "facts": facts_out,
        "examples": examples_out,
    }


def main() -> None:
    if len(sys.argv) < 3:
        print(
            "Usage: python extract_excel.py <featureId> <path/to/file.xlsx>",
            file=sys.stderr,
        )
        sys.exit(1)

    fid, xlsx_path = sys.argv[1], sys.argv[2]
    _log.info("Loading '%s'  (featureId=%s)", xlsx_path, fid)
    wb = openpyxl.load_workbook(xlsx_path, data_only=True, read_only=True)

    raw = xlsx_path.replace("\\", "/").split("/")[-1]
    stem = sanitize(raw[:-5] if raw.lower().endswith(".xlsx") else raw)

    results = [
        e
        for name in wb.sheetnames
        if (ws := wb[name]) is not None
        and (e := process_sheet(ws, name, stem, fid))
        and (e["facts"] or e["examples"])
    ]

    _log.info("Done: %d extraction(s) emitted to stdout", len(results))
    sys.stdout.buffer.write(
        json.dumps(results, ensure_ascii=False, indent=2).encode("utf-8")
    )
    sys.stdout.buffer.flush()


if __name__ == "__main__":
    main()
